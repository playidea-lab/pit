"""월간 리포트 조립: 수집 → 엑셀(일별·월별) → 메일 본문(월별 요약만)."""

import html
import logging
from dataclasses import dataclass
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from git_watcher.calendar_kr import is_workday, month_days
from git_watcher.config import Settings
from git_watcher.gitlab import GitLabClient
from git_watcher.identity import build_index
from git_watcher.monthly import (
    PersonMonth,
    add_commits,
    add_events,
    daily_rows,
    estimate_days,
    month_window,
    narrate,
    summary_cells,
)
from git_watcher.people import load_people
from git_watcher.sources import collect_all
from git_watcher.summarize import Summarizer

logger = logging.getLogger(__name__)

DAILY_HEADERS = ["날짜", "요일", "구분", "이름", "첫 활동", "마지막 활동", "활동 폭(시간)",
                 "수작업 환산(하한)", "수작업 환산(상한)", "하루 8시간 대비(중간값)", "커밋", "추가 줄", "삭제 줄", "기타 활동(push·이슈·MR·댓글)"]
NOTE = ("시각은 근무 시각이 아니라 GitLab에 남은 흔적(커밋 작성·push·이슈·MR·댓글)의 첫~마지막 시각입니다. "
        "첫 흔적 이전과 마지막 흔적 이후의 작업·회의, 커밋 없이 한 일은 보이지 않습니다. "
        "평균은 활동이 있었던 근무일(주말·공휴일 제외)만으로 계산했습니다. "
        "수작업 환산 작업량은 날마다 커밋 결과물을 'AI 없이 숙련 개발자가 손으로 만들었다면 걸렸을 시간'으로 "
        "AI가 추정한 범위의 합이며(주말·공휴일 포함), 실제로 일한 시간이 아니라 결과물의 크기입니다.")
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
OFFDAY_FILL = PatternFill("solid", fgColor="FBEFEF")
MAX_COLUMN_WIDTH = 60


@dataclass
class MonthlyReport:
    year: int
    month: int
    people: list[PersonMonth]
    workday_total: int

    @property
    def label(self) -> str:
        return f"{self.year}년 {self.month}월"

    @property
    def subject(self) -> str:
        return f"[GitLab 월간 리포트] {self.label} — 일별 활동 시각·월별 요약"

    @property
    def filename(self) -> str:
        return f"gitlab-activity-{self.year}-{self.month:02d}.xlsx"


def build_monthly(settings: Settings, year: int, month: int, summarizer: Summarizer | None) -> MonthlyReport:
    start, end = month_window(year, month, ZoneInfo(settings.timezone))
    names, exclude = load_people()
    with GitLabClient(settings.gitlab_url, settings.gitlab_token.get_secret_value()) as gl:
        # 이메일 매핑은 제외 계정까지 포함한 전원으로 만든다 (대표 커밋이 남에게 귀속되지 않도록)
        all_users = gl.human_users()
        index = build_index(gl, all_users, start, settings.state_dir)
        accounts = collect_all(settings, gl, start, end, index)  # 일일 추정과 같은 입력이 되도록 파일 경로 포함
        users = [u for u in all_users if u["username"] not in exclude]
        people = []
        for user in sorted(users, key=lambda u: names.get(u["username"], u["name"])):
            person = PersonMonth(user["username"], names.get(user["username"], user["name"]))
            if user["username"] in accounts:
                add_commits(person, accounts[user["username"]], start.tzinfo)
            add_events(person, gl, user["id"], start, end)
            people.append(person)
    if summarizer:
        estimate_days(summarizer, people, start.tzinfo)
    for person in people:
        person.narrative = narrate(summarizer, person) if summarizer else None
    workdays = sum(1 for d in month_days(year, month) if is_workday(d))
    return MonthlyReport(year, month, people, workdays)


def write_xlsx(report: MonthlyReport, out_dir: Path) -> Path:
    wb = Workbook()
    summary = wb.active
    summary.title = "월별 요약"
    cells = [summary_cells(p, report.workday_total) for p in report.people]
    headers = [*cells[0].keys(), "이번 달 한 일", "주요 작업"] if cells else []
    summary.append(headers)
    for person, row in zip(report.people, cells, strict=True):
        nar = person.narrative
        summary.append([*row.values(), nar.headline if nar else "", "\n".join(nar.items) if nar else ""])
    summary.append([])
    summary.append([NOTE])

    daily = wb.create_sheet("일별 활동")
    daily.append(DAILY_HEADERS)
    for person in report.people:
        for row in daily_rows(person, report.year, report.month):
            daily.append(row)
            if row[2] != "근무일":
                for cell in daily[daily.max_row]:
                    cell.fill = OFFDAY_FILL
    for sheet in (summary, daily):
        style_sheet(sheet)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / report.filename
    wb.save(path)
    return path


def style_sheet(sheet) -> None:  # noqa: ANN001 — openpyxl Worksheet
    """머리글 강조·틀 고정·자동 필터·열 너비."""
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for col in sheet.columns:
        width = max(len(str(c.value or "").split("\n")[0]) for c in col if c.row != sheet.max_row)
        sheet.column_dimensions[col[0].column_letter].width = min(width * 1.6 + 2, MAX_COLUMN_WIDTH)
        for c in col:
            c.alignment = Alignment(wrap_text=True, vertical="top")


def to_html(report: MonthlyReport) -> str:
    cell = "padding:6px 8px;border-top:1px solid #e5e5e5;vertical-align:top"
    rows = [summary_cells(p, report.workday_total) for p in report.people]
    head = "".join(f"<th style='text-align:left;padding:6px 8px'>{h}</th>" for h in rows[0]) if rows else ""
    body = "".join(
        "<tr>" + "".join(f"<td style='{cell};white-space:nowrap'>{html.escape(str(v))}</td>" for v in r.values())
        + "</tr>" for r in rows
    )
    details = "".join(_narrative_html(p) for p in report.people)
    return ("<div style='font-family:sans-serif;font-size:14px;color:#222'>"
            f"<h2>{html.escape(report.label)} 월별 요약</h2>"
            f"<p style='color:#666'>근무일 {report.workday_total}일 (주말·공휴일 제외). "
            "날짜별 첫·마지막 활동 시각은 첨부한 엑셀의 '일별 활동' 시트에 있습니다.</p>"
            "<table style='border-collapse:collapse;width:100%'>"
            f"<tr style='background:#f2f2f2'>{head}</tr>{body}</table>{details}"
            f"<p style='color:#999;font-size:12px;margin-top:24px'>{html.escape(NOTE)}</p></div>")


def _narrative_html(person: PersonMonth) -> str:
    if not person.narrative:
        return ""
    items = "".join(f"<li>{html.escape(i)}</li>" for i in person.narrative.items)
    return (f"<h3 style='margin:20px 0 4px'>{html.escape(person.name)}</h3>"
            f"<div>{html.escape(person.narrative.headline)}</div>"
            f"<ul style='line-height:1.7;margin-top:6px'>{items}</ul>")


def to_text(report: MonthlyReport) -> str:
    lines = [f"{report.label} 월별 요약 (근무일 {report.workday_total}일)", ""]
    for person in report.people:
        row = summary_cells(person, report.workday_total)
        lines.append(" · ".join(f"{k} {v}" for k, v in row.items()))
        if person.narrative:
            lines += [f"  {person.narrative.headline}", *[f"  - {i}" for i in person.narrative.items]]
    return "\n".join([*lines, "", NOTE, "", "일별 활동 시각은 첨부한 엑셀을 보세요."])
