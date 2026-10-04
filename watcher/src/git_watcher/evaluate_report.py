"""평가 조립: 수집 → 사람별 채점 → 엑셀(종합·근거·기준) + 메일 요약."""

import html
import logging
from dataclasses import dataclass
from datetime import timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import httpx
from openpyxl import Workbook

from git_watcher.config import Settings
from git_watcher.evaluate import (
    MIN_COMMITS_FOR_SCORE,
    RUBRIC,
    SECRET_PATH,
    ItemResult,
    SecretFinding,
    measure_all,
    scored_weight,
    secret_findings,
    total_score,
)
from git_watcher.gitlab import Commit, GitLabClient
from git_watcher.identity import build_index
from git_watcher.monthly import month_window
from git_watcher.monthly_report import style_sheet
from git_watcher.people import load_people
from git_watcher.sources import collect_all, commit_diff

logger = logging.getLogger(__name__)

INSUFFICIENT = "표본 부족"
NOTE = ("평가 v1은 PI Lab 규칙 파일에서 git 기록으로 검증할 수 있는 항목만 채점합니다. "
        "연구 저장소(research/)는 직접 push·테스트 동반 항목에서 면제합니다. "
        "에이전트 표기는 Claude만 남기고 Codex는 남기지 않아 참고로만 보여 주고 점수에는 넣지 않습니다. "
        "'채점된 가중치'가 작으면 면제·대상 없음 항목이 많아 총점의 근거가 좁다는 뜻이니 사람 간 비교에 주의하세요. "
        "에이전트에게 무엇을 지시했는지, 리뷰를 얼마나 성실히 했는지는 세션 기록이 연결되기 전까지 보이지 않습니다. "
        f"커밋 {MIN_COMMITS_FOR_SCORE}건 미만이면 총점을 내지 않습니다.")


@dataclass
class PersonEval:
    name: str
    commits: list[Commit]
    items: dict[str, ItemResult]

    @property
    def total(self) -> float | None:
        return total_score(self.items, len(self.commits))

    @property
    def coverage(self) -> str:
        done, full = scored_weight(self.items)
        return f"{done}/{full}"

    @property
    def total_label(self) -> str:
        return INSUFFICIENT if self.total is None else f"{self.total:g}"


def project_info(gl: GitLabClient, project_ids: set[int]) -> dict[int, dict]:
    """프로젝트 id → {default_branch, path}. 삭제된 프로젝트는 빠진다 (직접 push 판정 제외)."""
    info = {}
    for pid in project_ids:
        try:
            p = gl.project(pid)
        except httpx.HTTPStatusError as e:
            logger.warning("프로젝트 조회 실패 %s: %s", pid, e)
            continue
        info[pid] = {"default_branch": p.get("default_branch") or "", "path": p["path_with_namespace"]}
    return info


def scan_secrets(settings: Settings, gl: GitLabClient, commits: list[Commit]) -> dict[str, list[SecretFinding]]:
    """env류 파일을 건드린 커밋만 diff를 열어 비밀값 변수 이름을 찾는다."""
    findings: dict[str, list[SecretFinding]] = {}
    for c in commits:
        if not any(SECRET_PATH.search(f) for f in c.files):
            continue
        for d in commit_diff(settings, gl, c.project, c.sha, c.source):
            findings.setdefault(c.sha, []).extend(
                secret_findings(d["new_path"], d.get("diff", ""), d.get("deleted_file", False)))
    return findings


def build_evaluation(settings: Settings, year: int, month: int) -> list[PersonEval]:
    start, end = month_window(year, month, ZoneInfo(settings.timezone))
    names, exclude = load_people()
    with GitLabClient(settings.gitlab_url, settings.gitlab_token.get_secret_value()) as gl:
        all_users = gl.human_users()
        accounts = collect_all(settings, gl, start, end, build_index(gl, all_users, start, settings.state_dir))
        users = [u for u in all_users if u["username"] not in exclude]
        after, before = (start.date() - timedelta(days=1)).isoformat(), end.date().isoformat()
        pushes = {u["username"]: [e for e in gl.events(u["id"], after, before)
                                  if (e.get("action_name") or "").startswith("pushed")] for u in users}
        projects = project_info(gl, {e["project_id"] for es in pushes.values() for e in es})
        secrets = scan_secrets(settings, gl, [c for a in accounts.values() for c in a.commits])
    result = []
    for user in users:
        commits = accounts[user["username"]].commits if user["username"] in accounts else []
        items = measure_all(commits, pushes[user["username"]], projects, secrets,
                            tuple(settings.research_namespaces))
        result.append(PersonEval(names.get(user["username"], user["name"]), commits, items))
    return sorted(result, key=lambda p: -(p.total or -1))


def write_xlsx(people: list[PersonEval], label: str, out_dir: Path) -> Path:
    wb = Workbook()
    summary = wb.active
    summary.title = "종합"
    summary.append(["이름", "총점(100)", "채점된 가중치", "커밋", *[f"{c.title} ({c.weight})" if c.weight else c.title for c in RUBRIC]])
    for p in people:
        summary.append([p.name, p.total_label, p.coverage, len(p.commits), *[_cell(p.items[c.key]) for c in RUBRIC]])
    summary.append([])
    summary.append([NOTE])

    evidence = wb.create_sheet("항목별 근거")
    evidence.append(["이름", "항목", "점수", "통과/대상", "위반 예시 (최대 5건)"])
    for p in people:
        for c in RUBRIC:
            item = p.items[c.key]
            count = f"위반 {item.total}건" if item.fixed_score is not None else f"{item.passed}/{item.total}"
            evidence.append([p.name, c.title, _cell(item), count, "\n".join(item.violations)])

    rubric = wb.create_sheet("채점 기준")
    rubric.append(["분류", "항목", "가중치", "근거 규칙", "측정 방법", "한계"])
    for c in RUBRIC:
        rubric.append([c.category, c.title, c.weight, c.rule, c.how, c.caveat])
    for sheet in (summary, evidence, rubric):
        style_sheet(sheet)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"evaluation-v1-{label}.xlsx"
    wb.save(path)
    return path


def _cell(item: ItemResult) -> str:
    return "-" if item.score is None else f"{item.score:g}"


def to_html(people: list[PersonEval], label: str) -> str:
    td = "padding:6px 8px;border-top:1px solid #e5e5e5;vertical-align:top;white-space:nowrap"
    head = "".join(f"<th style='text-align:left;padding:6px 8px'>{html.escape(h)}</th>"
                   for h in ["이름", "총점", "채점된 가중치", "커밋", *[c.title for c in RUBRIC]])
    rows = "".join(
        "<tr>" + "".join(f"<td style='{td}'>{html.escape(str(v))}</td>"
                         for v in [p.name, p.total_label, p.coverage, len(p.commits), *[_cell(p.items[c.key]) for c in RUBRIC]])
        + "</tr>" for p in people)
    weights = " · ".join(f"{c.title} {c.weight}" for c in RUBRIC)
    return ("<div style='font-family:sans-serif;font-size:14px;color:#222'>"
            f"<h2>작업 방식 평가 v1 — {html.escape(label)}</h2>"
            f"<p style='color:#666'>가중치: {html.escape(weights)}. 항목별 위반 근거(커밋 링크)는 첨부 엑셀에 있습니다.</p>"
            f"<table style='border-collapse:collapse'><tr style='background:#f2f2f2'>{head}</tr>{rows}</table>"
            f"<p style='color:#999;font-size:12px;margin-top:24px'>{html.escape(NOTE)}</p></div>")


def to_text(people: list[PersonEval], label: str) -> str:
    lines = [f"작업 방식 평가 v1 — {label}", ""]
    for p in people:
        parts = " · ".join(f"{c.title} {_cell(p.items[c.key])}" for c in RUBRIC)
        lines.append(f"{p.name}: 총점 {p.total_label} (채점된 가중치 {p.coverage}, 커밋 {len(p.commits)}) — {parts}")
    return "\n".join([*lines, "", NOTE])
